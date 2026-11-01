"""Builds DMS_VIDEO_CATALOGUE.xlsx from DMS_VIDEO_CATALOGUE.csv (the csv is the master; run after adding a row)."""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

here = Path(__file__).parent
rows = list(csv.reader(open(here / "DMS_VIDEO_CATALOGUE.csv", encoding="utf8")))
wb = Workbook()
ws = wb.active
ws.title = "Videos"
for r in rows:
    ws.append(r)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F4E79")
for col, w in zip("ABCDEFGH", (5, 40, 30, 10, 8, 90, 48, 24)):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"
wb.save(here / "DMS_VIDEO_CATALOGUE.xlsx")
